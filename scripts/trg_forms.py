#!/usr/bin/env python3
"""Buat formulir kelengkapan data per TRG (xlsx) yang sudah terisi dari data crawl.

Pemakaian:  python3 scripts/trg_forms.py [--out forms/trg]
Satu berkas per TRG: Petunjuk, Anggota, TRG & Riset, Daftar (pilihan dropdown).
Sel abu-abu = sudah terisi dari situs TRG (periksa dan koreksi); sel kuning = perlu diisi.
Hasil isian dibaca kembali dengan scripts/trg_forms_collect.py.
"""
import argparse, csv, os, collections as C
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

D = "data/trg"
R = lambda f: list(csv.DictReader(open(f"{D}/{f}.csv", encoding="utf-8")))
trg, pen, tags, tt, lap = R("trg"), R("peneliti"), R("tag_tier1"), R("trg_tag_tier1"), R("laporan_kelengkapan")
tag_names = [t["nama_id"] for t in tags]
tname = {t["tag_id"]: t["nama_id"] for t in tags}
tmap = {r["slug"]: [tname[x] for x in r["tag_ids"].split(";") if x] for r in tt}
lapd = {r["slug"]: r for r in lap}
by = C.defaultdict(list)
for p in pen:
    by[p["trg_slug"]].append(p)

JENIS_MITRA = ["Pemerintah pusat", "Pemerintah daerah", "Industri/swasta", "LSM/komunitas", "Perguruan tinggi dalam negeri",
               "Perguruan tinggi/lembaga luar negeri", "Lembaga riset (BRIN dll.)", "Rumah sakit/fasilitas kesehatan", "Lainnya"]
STATUS = ["Aktif", "Tidak aktif lagi"]
TKT = [str(i) for i in range(1, 10)]
LUARAN = ["Artikel jurnal", "Buku/bab buku", "HKI/paten", "Produk/prototipe", "Kebijakan/rekomendasi", "Pelatihan/pendampingan", "Lainnya"]
EXTRA_ROWS = 12

F = "Arial"
f_norm, f_bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True)
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_note = Font(name=F, size=9, italic=True, color="595959")
fill_head = PatternFill("solid", fgColor="1F3864")
fill_fill = PatternFill("solid", fgColor="FFF2CC")   # kuning: isi
fill_pre = PatternFill("solid", fgColor="E7E6E6")    # abu: terisi dari situs, periksa
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")


def put(ws, ref, v, font=f_norm, fill=None, align=wrap, border=True):
    c = ws[ref]
    c.value = v; c.font = font; c.alignment = align
    if fill: c.fill = fill
    if border: c.border = box
    return c


def header(ws, row, labels, widths=None):
    for i, l in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=l)
        c.font, c.fill, c.border = f_head, fill_head, box
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 32


def build(t, path):
    slug = t["slug"]; name = t["nama"]
    members = by.get(slug, [])
    l = lapd[slug]
    masalah = [m for m in l["masalah"].split(";") if m]
    wb = Workbook()

    # ---- Daftar ----
    wd = wb.active; wd.title = "Daftar"
    lists = [("Tema", tag_names), ("Jenis mitra", JENIS_MITRA), ("Status", STATUS), ("TKT", TKT), ("Luaran", LUARAN)]
    for j, (h, vals) in enumerate(lists, 1):
        wd.cell(row=1, column=j, value=h).font = f_bold
        wd.column_dimensions[get_column_letter(j)].width = 44 if j in (1, 2) else 22
        for i, v in enumerate(vals, 2):
            wd.cell(row=i, column=j, value=v).font = f_norm
    rng = lambda col, n: f"=Daftar!${col}$2:${col}${n + 1}"
    dv_tema = DataValidation(type="list", formula1=rng("A", len(tag_names)), allow_blank=True)
    dv_jenis = DataValidation(type="list", formula1=rng("B", len(JENIS_MITRA)), allow_blank=True)
    dv_status = DataValidation(type="list", formula1=rng("C", len(STATUS)), allow_blank=True)
    dv_tkt = DataValidation(type="list", formula1=rng("D", len(TKT)), allow_blank=True)
    dv_luaran = DataValidation(type="list", formula1=rng("E", len(LUARAN)), allow_blank=True)

    # ---- Anggota ----
    wa = wb.create_sheet("Anggota", 0)
    wa["A1"].value = f"Anggota {name}"; wa["A1"].font = f_title
    wa["A2"].value = ("Abu-abu: sudah terisi dari situs TRG, mohon diperiksa dan dikoreksi. Kuning: mohon diisi. "
                      "Tambahkan anggota yang belum tercantum pada baris kosong di bawah.")
    wa["A2"].font = f_note
    wa.merge_cells("A2:M2"); wa["A2"].alignment = wrap; wa.row_dimensions[2].height = 28
    cols = ["No", "Nama lengkap dan gelar", "Peran (Leader/Lecturer/Researcher/Student/External)", "Afiliasi",
            "Keahlian di situs TRG", "Keahlian utama 1 (pilih)", "Keahlian utama 2 (pilih)", "Keahlian utama 3 (pilih)",
            "Kata kunci keahlian (bebas, pisahkan koma)", "ORCID", "SINTA ID", "Scopus ID", "Status"]
    header(wa, 4, cols, [5, 34, 22, 30, 30, 30, 30, 30, 34, 22, 12, 16, 14])
    start = 5
    nrows = max(len(members), 0) + EXTRA_ROWS
    for i in range(nrows):
        r = start + i
        m = members[i] if i < len(members) else None
        vals = [i + 1] + ([m["nama"], m["peran"], m["afiliasi"], m["keahlian"], "", "", "", "", m["orcid"], m["sinta_id"], m["scopus_id"], "Aktif"]
                          if m else [""] * 12)
        for j, v in enumerate(vals, 1):
            c = wa.cell(row=r, column=j, value=v if v != "" else None)
            c.font, c.border, c.alignment = f_norm, box, wrap
            if j == 1:
                continue
            c.fill = fill_pre if (m and v not in ("", None) and j in (2, 3, 4, 5, 10, 11, 12, 13)) else fill_fill
            if j == 5:
                c.fill = fill_pre
    last = start + nrows - 1
    for dv, col in ((dv_tema, "F"), (dv_tema, "G"), (dv_tema, "H"), (dv_status, "M")):
        pass
    dv_a = DataValidation(type="list", formula1=rng("A", len(tag_names)), allow_blank=True)
    dv_s = DataValidation(type="list", formula1=rng("C", len(STATUS)), allow_blank=True)
    dv_a.add(f"F{start}:H{last}"); dv_s.add(f"M{start}:M{last}")
    wa.add_data_validation(dv_a); wa.add_data_validation(dv_s)
    wa.freeze_panes = "C5"
    wa.sheet_view.zoomScale = 90

    # ---- TRG & Riset ----
    wt = wb.create_sheet("TRG & Riset", 1)
    for col, w in zip("ABCDEFG", (34, 40, 30, 24, 12, 26, 36)):
        wt.column_dimensions[col].width = w
    wt["A1"].value = f"Profil {name}"; wt["A1"].font = f_title
    put(wt, "A3", "Fakultas", f_bold); put(wt, "B3", t["fakultas"].replace("Faculty of ", "").replace("School of ", ""), f_norm, fill_pre)
    put(wt, "A5", "Tema riset utama (pilih 1-3)", f_bold)
    draft = tmap.get(slug, [])[:3] + [""] * 3
    for k in range(3):
        put(wt, f"B{5 + k}", draft[k] or None, f_norm, fill_pre if draft[k] else fill_fill)
    put(wt, "C5", "Tema terisi berdasarkan nama TRG (draf otomatis). Ganti bila kurang tepat.", f_note, border=False)
    dv_t = DataValidation(type="list", formula1=rng("A", len(tag_names)), allow_blank=True)
    dv_t.add("B5:B7"); wt.add_data_validation(dv_t)
    put(wt, "A9", "Kata kunci riset (5-10, pisahkan koma)", f_bold); put(wt, "B9", None, f_norm, fill_fill)
    put(wt, "A10", "Deskripsi singkat TRG (maks. 3 kalimat)", f_bold); put(wt, "B10", None, f_norm, fill_fill)
    wt.row_dimensions[9].height = 42; wt.row_dimensions[10].height = 64
    wt.merge_cells("B9:G9"); wt.merge_cells("B10:G10")

    put(wt, "A12", "Mitra kerja sama (maks. 5)", f_bold, border=False)
    for j, h in enumerate(["Nama mitra", "Jenis mitra (pilih)", "Bentuk kerja sama (hibah, riset bersama, magang, dll.)", "Tahun mulai"], 1):
        c = wt.cell(row=13, column=j, value=h); c.font, c.fill, c.border, c.alignment = f_head, fill_head, box, wrap
    for r in range(14, 19):
        for j in range(1, 5):
            c = wt.cell(row=r, column=j); c.fill, c.border, c.font, c.alignment = fill_fill, box, f_norm, wrap
    dv_j = DataValidation(type="list", formula1=rng("B", len(JENIS_MITRA)), allow_blank=True); dv_j.add("B14:B18"); wt.add_data_validation(dv_j)

    put(wt, "A20", "Riset terbaik 5 tahun terakhir (maks. 3)", f_bold, border=False)
    for j, h in enumerate(["Judul riset", "Sumber dana", "Mitra", "Tahun", "TKT (1-9)", "Luaran utama (pilih)", "Tautan (DOI/URL)"], 1):
        c = wt.cell(row=21, column=j, value=h); c.font, c.fill, c.border, c.alignment = f_head, fill_head, box, wrap
    for r in range(22, 25):
        for j in range(1, 8):
            c = wt.cell(row=r, column=j); c.fill, c.border, c.font, c.alignment = fill_fill, box, f_norm, wrap
    dv_k = DataValidation(type="list", formula1=rng("D", len(TKT)), allow_blank=True); dv_k.add("E22:E24"); wt.add_data_validation(dv_k)
    dv_l = DataValidation(type="list", formula1=rng("E", len(LUARAN)), allow_blank=True); dv_l.add("F22:F24"); wt.add_data_validation(dv_l)
    put(wt, "A26", "TKT = Tingkat Kesiapterapan Teknologi. Isi bila riset menghasilkan produk atau teknologi, kosongkan untuk riset dasar.", f_note, border=False)

    # ---- Petunjuk ----
    wp = wb.create_sheet("Petunjuk", 0)
    wp.column_dimensions["A"].width = 42; wp.column_dimensions["B"].width = 70
    wp["A1"].value = f"Formulir kelengkapan data: {name}"; wp["A1"].font = f_title
    wp["A2"].value = "Peta Riset Unhas, Thematic Research Group"; wp["A2"].font = f_note
    rows = [
        ("Tujuan", "Melengkapi data agar peneliti TRG ini dapat ditemukan di dashboard Peta Riset Unhas (siapa ahli tentang apa)."),
        ("Lama pengisian", "Sekitar 15 menit."),
        ("Sel abu-abu", "Sudah terisi dari situs TRG. Mohon diperiksa dan dikoreksi bila keliru."),
        ("Sel kuning", "Mohon diisi. Boleh dikosongkan bila tidak ada."),
        ("Lembar Anggota", "Periksa daftar anggota. Pilih sampai 3 keahlian utama dari daftar, lalu tulis kata kunci bebas. Isi ORCID/SINTA/Scopus bila punya."),
        ("Lembar TRG & Riset", "Pilih tema utama, tulis kata kunci dan deskripsi singkat, sebutkan mitra dan riset terbaik."),
        ("Keahlian ditampilkan publik", "Keahlian, afiliasi, dan ID peneliti akan tampil di dashboard Peta Riset Unhas."),
        ("Kembalikan ke", "trg@unhas.ac.id, dengan berkas ini terlampir."),
    ]
    for i, (a, b) in enumerate(rows, 4):
        put(wp, f"A{i}", a, f_bold); put(wp, f"B{i}", b)
    r0 = 4 + len(rows) + 1
    put(wp, f"A{r0}", "Yang paling perlu dilengkapi", f_bold, border=False)
    need = []
    if "tanpa_halaman_situs" in masalah:
        need.append("TRG ini belum punya halaman sendiri di situs TRG. Mohon isi seluruh formulir, termasuk daftar anggota.")
    if "halaman_tim_404" in masalah or "tim_tidak_terbaca_(tata_letak_lain_atau_kosong)" in masalah:
        need.append("Daftar anggota tim belum terbaca dari situs. Mohon isi lengkap di lembar Anggota.")
    if "keahlian_<50%" in masalah:
        need.append("Keahlian anggota banyak yang kosong. Mohon pilih keahlian utama untuk setiap anggota.")
    if "ID_peneliti_<30%" in masalah:
        need.append("ID peneliti (SINTA/Scopus/ORCID) banyak yang kosong.")
    if "tanpa_publikasi" in masalah:
        need.append("Belum ada artikel tercatat di situs. Mohon cantumkan riset terbaik di lembar TRG & Riset.")
    if "tanpa_proyek" in masalah:
        need.append("Belum ada proyek riset tercatat di situs.")
    need = need or ["Data TRG ini relatif lengkap. Mohon periksa dan lengkapi mitra serta riset terbaik."]
    for i, n in enumerate(need, r0 + 1):
        put(wp, f"A{i}", f"{i - r0}.", f_norm); put(wp, f"B{i}", n)
    r1 = r0 + len(need) + 2
    put(wp, f"A{r1}", "Contoh pengisian (hanya ilustrasi, jangan disalin)", f_bold, border=False)
    ex = [("Nama", "Dr. Contoh Nama, S.Pi., M.Si."), ("Keahlian utama 1-3 (pilih)", "Perikanan & akuakultur; Ekosistem pesisir & laut"),
          ("Kata kunci keahlian", "budidaya udang, kualitas air, bioflok"), ("SINTA ID", "6012345"), ("ORCID", "0000-0002-1234-5678")]
    for i, (a, b) in enumerate(ex, r1 + 1):
        put(wp, f"A{i}", a, f_norm); put(wp, f"B{i}", b, f_note)
    r2 = r1 + len(ex) + 2
    put(wp, f"A{r2}", "Kemajuan pengisian (otomatis)", f_bold, border=False)
    put(wp, f"A{r2 + 1}", "Anggota tercantum"); put(wp, f"B{r2 + 1}", f"=COUNTA(Anggota!B{start}:B{last})", align=Alignment(horizontal="left"))
    put(wp, f"A{r2 + 2}", "Anggota dengan keahlian terpilih"); put(wp, f"B{r2 + 2}", f"=COUNTA(Anggota!F{start}:F{last})", align=Alignment(horizontal="left"))
    put(wp, f"A{r2 + 3}", "Persentase anggota berkeahlian"); put(wp, f"B{r2 + 3}", f"=IF(B{r2 + 1}=0,0,B{r2 + 2}/B{r2 + 1})", align=Alignment(horizontal="left"))
    wp[f"B{r2 + 3}"].number_format = "0%"
    wb.move_sheet("Daftar", offset=0)
    wb._sheets = [wb["Petunjuk"], wb["Anggota"], wb["TRG & Riset"], wb["Daftar"]]
    wb.active = 0
    wb.save(path)
    return len(members), need


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", default="forms/trg"); a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    idx = []
    for t in trg:
        fn = f"Formulir_TRG_{t['slug']}.xlsx"
        n, need = build(t, os.path.join(a.out, fn))
        idx.append([t["slug"], t["nama"], t["fakultas"], fn, n, len(lapd[t["slug"]]["masalah"].split(";")) if lapd[t["slug"]]["masalah"] else 0, " | ".join(need)])
    idx.sort(key=lambda r: -r[5])
    with open(os.path.join(os.path.dirname(a.out.rstrip("/")) or ".", "index_formulir.csv"), "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh); w.writerow(["slug", "nama", "fakultas", "berkas", "n_anggota_terisi", "n_masalah", "prioritas_isi"]); w.writerows(idx)
    print(len(idx), "formulir dibuat di", a.out)


if __name__ == "__main__":
    main()
