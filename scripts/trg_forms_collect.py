#!/usr/bin/env python3
"""Baca kembali formulir TRG yang sudah diisi (xlsx) menjadi CSV.

Pemakaian: python3 scripts/trg_forms_collect.py DIR_FORMULIR_MASUK [--out data/trg/formulir_masuk]
Keluaran: anggota.csv, trg_profil.csv, mitra.csv, riset.csv. Nama berkas harus Formulir_TRG_<slug>.xlsx.
"""
import argparse, csv, glob, os, re
from openpyxl import load_workbook

ap = argparse.ArgumentParser(); ap.add_argument("dir"); ap.add_argument("--out", default="data/trg/formulir_masuk"); a = ap.parse_args()
os.makedirs(a.out, exist_ok=True)
cell = lambda v: ("" if v is None else str(v).strip())
anggota, profil, mitra, riset, bad = [], [], [], [], []
for f in sorted(glob.glob(os.path.join(a.dir, "Formulir_TRG_*.xlsx"))):
    slug = re.sub(r"^Formulir_TRG_|\.xlsx$", "", os.path.basename(f))
    try:
        wb = load_workbook(f, data_only=True)
        wa, wt = wb["Anggota"], wb["TRG & Riset"]
    except Exception as e:
        bad.append((slug, str(e))); continue
    for r in wa.iter_rows(min_row=5, values_only=True):
        if cell(r[1]):
            anggota.append([slug] + [cell(x) for x in r[1:13]])
    profil.append([slug, "; ".join(x for x in (cell(wt[f"B{i}"].value) for i in (5, 6, 7)) if x), cell(wt["B9"].value), cell(wt["B10"].value)])
    for r in wt.iter_rows(min_row=14, max_row=18, max_col=4, values_only=True):
        if cell(r[0]): mitra.append([slug] + [cell(x) for x in r])
    for r in wt.iter_rows(min_row=22, max_row=24, max_col=7, values_only=True):
        if cell(r[0]): riset.append([slug] + [cell(x) for x in r])

def write(name, head, rows):
    with open(os.path.join(a.out, name), "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh); w.writerow(head); w.writerows(rows)
write("anggota.csv", ["trg_slug", "nama", "peran", "afiliasi", "keahlian_situs", "keahlian_1", "keahlian_2", "keahlian_3", "kata_kunci", "orcid", "sinta_id", "scopus_id", "status"], anggota)
write("trg_profil.csv", ["trg_slug", "tema_utama", "kata_kunci", "deskripsi"], profil)
write("mitra.csv", ["trg_slug", "mitra", "jenis", "bentuk", "tahun_mulai"], mitra)
write("riset.csv", ["trg_slug", "judul", "sumber_dana", "mitra", "tahun", "tkt", "luaran", "tautan"], riset)
print(f"{len(profil)} formulir dibaca; {len(anggota)} anggota, {len(mitra)} mitra, {len(riset)} riset; gagal: {bad}")
