#!/usr/bin/env python3
"""Penyaringan kandidat konsorsium/institut dari data TRG -> Potensi_Konsorsium_TRG.xlsx.

Pemakaian: python3 scripts/trg_konsorsium.py OUT_XLSX
Ini penyaringan berbasis data publik TRG (ukuran, keterhubungan, mitra luar), bukan keputusan kelayakan:
pendanaan, tata kelola, komitmen ketua, dan status hukum tidak ada di data.
"""
import csv, itertools, re, sys, collections as C
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

D = "data/trg"
R = lambda f: list(csv.DictReader(open(f"{D}/{f}.csv", encoding="utf-8")))
trg = {t["slug"]: t for t in R("trg")}
pen, unik, pub = R("peneliti"), R("peneliti_unik"), R("publikasi")
roles = lambda p: {x.strip().lower() for x in p["peran"].split(";") if x.strip()}
norm = lambda t: re.sub(r"[^a-z0-9]+", " ", t.lower()).strip()
arts = {}
for r in pub:
    k = norm(r["judul"])
    if len(k) >= 12:
        arts.setdefault(k, set()).add(r["trg_slug"])
ART_PER = C.Counter(s for v in arts.values() for s in v)
INA = re.compile(r"indonesia|universitas|institut\b|politeknik|unhas|hasanuddin|makassar|brin|bogor|bandung|jakarta|surabaya|yogya|manado|sulawesi|kementerian|stikes|uin\b|ipb|itb|ugm|unm\b|unpad|undip|unair|bone|gowa|palu|kendari|ambon|mataram|papua|aceh|medan|padang|riau|pontianak|banjarmasin|malang|semarang|udayana", re.I)

# Kandidat: aturan kata kunci pada nama TRG (kelompok buatan saya; satu TRG boleh masuk beberapa kandidat)
# (nama, regex, bentuk_disarankan, pusat_unggulan_terkait, catatan)
CAND = [
 ("Kemaritiman, pesisir & perikanan", r"coastal|marine|coral|seagrass|mangrove|fisher|aquacult|shrimp|crab|lobster|tuna|shark|seaweed|algal|\bfish|blue econom|maritime|ocean|tilapia|aquatic|multitrophic|invasive alien|natural coastal|smart ship",
  "Institut (calon utama)", "Wallacea Research Center", "Paling lintas fakultas dan paling khas kawasan (Wallacea, Segitiga Terumbu Karang). Konsolidasikan dengan pusat yang sudah ada, jangan membuat unit baru yang menyaingi."),
 ("Penyakit tropis, infeksi & vaksin", r"tuberculosis|biovax|malaria|pneumonia|biofilm|superbug|microbiome|vector|infectious|vaccine|syndemic|public health entomology|unhas fly|animal disease",
  "Pusat riset/institut kompak", "(belum ada)", "Kecil tetapi padat: artikel per peneliti tertinggi dan mitra luar negeri nyata. Belum punya Pusat Unggulan."),
 ("Ketahanan pangan & pertanian tropis", r"rice|corn|food crop|sago|cocoa|coffee|oil palm|horticult|orchid|agricultur|soil|pest|phytophthora|mycolog|bacteriolog|biocontrol|biopesticide|post-?harvest|food|feed|livestock|dairy|poultry|beef|goat|equine|sub-optimal|land use|agro|taro|fermented|ethnic food",
  "Konsorsium, lalu pecah jadi 2-3 sub-konsorsium", "(belum ada)", "Terbesar tetapi terlalu luas untuk satu institut. Usul sub-konsorsium: perkebunan (kakao, kopi, sawit, sagu), tanaman pangan-hortikultura, peternakan."),
 ("Gizi, ibu-anak & penyakit tidak menular", r"nutrition|obesity|child|maternal|mother|neonat|birth cohort|ncd|non-communicable|diabetic|metabolism|nephro|cardiovascular|digestive|healthy aging|geriatric|pediatric|midwif",
  "Konsorsium (bisa naik jadi pusat)", "(belum ada)", "Catatan: Anda sendiri ketua salah satu TRG di kelompok ini, jadi nilai penilaian ini sebaiknya diperiksa pihak lain."),
 ("Tata kelola, perdamaian & kebijakan publik", r"peace|conflict|democracy|terror|migration|human security|governance|public organization|policy|constitution|government|public opinion|disab|criminolog|victim|legal protection|regulation|sociolegal|agrarian",
  "Perkuat Pusat Unggulan yang ada", "Pusat Perdamaian Global & Penyelesaian Konflik", "Ketua TRG paling banyak dan paling terhubung. Cocok menjadi lengan riset pusat yang sudah ada."),
 ("Kebencanaan & infrastruktur tangguh", r"earthquake|flood|flowslide|liquefaction|hazard|resilien|structural|geotechn|infrastructure|settlement|floating|piping|airport|transport|hydrology|watershed|climate based",
  "Konsorsium", "(belum ada)", "Keterhubungan antar-TRG dan kelengkapan ID peneliti rendah; mulai dari konsorsium, bukan institut."),
 ("Energi, material maju & mineral kritis", r"energy|hydrogen|renewable|power system|grid|waste to energy|solar|nanogenerator|semiconductor|material|nano|polymer|biochar|metallurg|rare-earth|rare earth|insulation|high voltage|graded|lightweight|fiber",
  "Perkuat Pusat Unggulan yang ada", "ASEAN-Indonesia-China CoE Metallurgy & Marine Resources", "Dua TRG logam kritis dan metalurgi berbagi peneliti paling banyak (10 orang). Gunakan sebagai inti."),
 ("Ekonomi, bisnis & akuntansi", r"account|audit|financ|econom|market|business|supply chain|sme\b|small and medium|human resource|strategic|corporate|budget|microfinance|accountab|transparen|forensic accounting|islamic",
  "Konsorsium", "(belum ada)", "Banyak TRG dan peneliti, tetapi satu rumpun fakultas; hasil riset lebih cocok sebagai layanan kajian ke pemerintah dan industri."),
 ("Bahan alam, farmasi & biomedis", r"natural product|drug|pharma|nutraceutical|cosmeceutical|metabolomic|bone|regenerat|dental|dentistry|oral|periodont|medicinal|transdermal|tumor|cancer|oncolog",
  "Konsorsium", "(belum ada)", "Tersebar di 8 fakultas dengan keterhubungan rendah; perlu pengikat tema yang jelas."),
]


def build():
    rows = []
    for name, rx, form, centre, note in CAND:
        rg = re.compile(rx, re.I)
        slugs = sorted((s for s, t in trg.items() if rg.search(t["nama"])), key=lambda s: -ART_PER[s])
        rows.append((name, form, centre, note, slugs))
    return rows


def main():
    out = sys.argv[1]
    cands = build()
    HF, HFILL, NF = Font(name="Arial", size=10, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864"), Font(name="Arial", size=10)
    wb = Workbook()
    ws = wb.active; ws.title = "Ringkasan"
    ws["A1"] = "Potensi konsorsium dan institut dari TRG Unhas (penyaringan awal)"; ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = ("Berbasis data publik TRG (ukuran, keterhubungan, mitra luar). Bukan keputusan kelayakan: pendanaan, tata kelola, komitmen ketua, "
                "dan status hukum tidak ada di data. Satu TRG dapat masuk lebih dari satu kandidat.")
    ws["A2"].font = Font(name="Arial", size=9, italic=True); ws.merge_cells("A2:L2"); ws["A2"].alignment = Alignment(wrap_text=True)
    head = ["Kandidat", "Bentuk yang disarankan", "Pusat Unggulan GloRI terkait", "Jumlah TRG", "Jumlah fakultas", "Peneliti (non-mahasiswa, unik)",
            "Ketua TRG terbaca", "Artikel unik", "Peneliti ber-ID (%)", "TRG saling terhubung (%)", "Mitra luar negeri (institusi)", "Catatan"]
    for j, h in enumerate(head, 1):
        c = ws.cell(row=4, column=j, value=h); c.font, c.fill, c.alignment = HF, HFILL, Alignment(wrap_text=True, vertical="center")
    detail = []
    for i, (name, form, centre, note, slugs) in enumerate(cands, 5):
        S = set(slugs)
        ppl = [p for p in unik if roles(p) != {"student"} and set(p["trg_slugs"].split(";")) & S]
        leads = [p for p in ppl if set(p["ketua_slugs"].split(";")) & S]
        art = {k for k, v in arts.items() if v & S}
        fak = {trg[s]["fakultas"] for s in slugs}
        linked = {s for a, b in itertools.combinations(slugs, 2) for p in ppl if {a, b} <= set(p["trg_slugs"].split(";")) for s in (a, b)}
        idp = sum(1 for p in ppl if p["scopus_id"] or p["orcid"] or p["sinta_id"])
        foreign = {re.sub(r"\s+", " ", p["afiliasi"]).strip().lower()[:60] for p in pen if p["trg_slug"] in S and p["peran"].lower().startswith("external")
                   and p["afiliasi"].strip() and not INA.search(p["afiliasi"]) and len(p["afiliasi"]) < 90}
        sheet = re.sub(r"[^A-Za-z0-9 ]", "", name)[:28].strip()
        vals = [name, form, centre, f"=COUNTA('{sheet}'!A5:A200)", len(fak), len(ppl), len(leads), len(art), round(idp / max(1, len(ppl)), 3),
                round(len(linked) / max(1, len(slugs)), 3), len(foreign), note]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v); c.font = NF; c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=9).number_format = "0%"; ws.cell(row=i, column=10).number_format = "0%"
        detail.append((sheet, name, slugs))
    for j, w in enumerate([36, 26, 30, 10, 10, 16, 12, 12, 12, 14, 14, 70], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[4].height = 44; ws.freeze_panes = "B5"
    for sheet, name, slugs in detail:
        w = wb.create_sheet(sheet)
        w["A1"] = name; w["A1"].font = Font(name="Arial", size=12, bold=True)
        w["A2"] = "TRG anggota kandidat (diurutkan menurut artikel unik)"; w["A2"].font = Font(name="Arial", size=9, italic=True)
        h = ["TRG", "Fakultas", "Ketua (yang terbaca)", "Peneliti (non-mahasiswa)", "Artikel unik", "Slug"]
        for j, x in enumerate(h, 1):
            c = w.cell(row=4, column=j, value=x); c.font, c.fill = HF, HFILL
        lead = {}
        for p in unik:
            for s in [x for x in p["ketua_slugs"].split(";") if x]:
                lead.setdefault(s, []).append(p["nama"])
        for r, s in enumerate(slugs, 5):
            n = sum(1 for p in unik if roles(p) != {"student"} and s in p["trg_slugs"].split(";"))
            vals = [trg[s]["nama"], trg[s]["fakultas"].replace("Faculty of ", ""), " | ".join(lead.get(s, [])) or "(belum terbaca)", n, ART_PER[s], s]
            for j, v in enumerate(vals, 1):
                w.cell(row=r, column=j, value=v).font = NF
        for j, wd in enumerate([52, 30, 44, 14, 12, 22], 1):
            w.column_dimensions[get_column_letter(j)].width = wd
        w.freeze_panes = "A5"
    # pasangan TRG yang paling banyak berbagi peneliti (kolaborasi nyata dari bawah)
    sh = C.Counter()
    for p in unik:
        if roles(p) == {"student"}:
            continue
        sl = sorted({s for s in p["trg_slugs"].split(";") if s in trg})
        for a, b in itertools.combinations(sl, 2):
            sh[(a, b)] += 1
    w = wb.create_sheet("Pasangan TRG terhubung")
    w["A1"] = "Pasangan TRG yang berbagi peneliti terbanyak (indikasi kolaborasi yang sudah berjalan)"; w["A1"].font = Font(name="Arial", size=12, bold=True)
    for j, x in enumerate(["TRG A", "TRG B", "Peneliti bersama", "Fakultas A", "Fakultas B"], 1):
        c = w.cell(row=3, column=j, value=x); c.font, c.fill = HF, HFILL
    for r, ((a, b), n) in enumerate(sh.most_common(40), 4):
        for j, v in enumerate([trg[a]["nama"], trg[b]["nama"], n, trg[a]["fakultas"].replace("Faculty of ", ""), trg[b]["fakultas"].replace("Faculty of ", "")], 1):
            w.cell(row=r, column=j, value=v).font = NF
    for j, wd in enumerate([50, 50, 16, 28, 28], 1):
        w.column_dimensions[get_column_letter(j)].width = wd
    wb.save(out)
    print("tersimpan", out, "| kandidat", len(cands))


if __name__ == "__main__":
    main()
