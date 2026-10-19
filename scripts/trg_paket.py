#!/usr/bin/env python3
"""Susun paket serah-terima: zip formulir, laporan kelengkapan (xlsx), dan data peta riset (xlsx).
Pemakaian: python3 scripts/trg_paket.py OUT_DIR   (jalankan setelah crawl, laporan, dedup, dan trg_forms)"""
import csv, glob, os, re, sys, zipfile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

P, D = sys.argv[1], "data/trg"
os.makedirs(P, exist_ok=True)
R = lambda f: list(csv.DictReader(open(f"{D}/{f}.csv", encoding="utf-8")))
HF, HFILL, NF = Font(name="Arial", size=10, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864"), Font(name="Arial", size=10)
NUM = re.compile(r"^(n_.*|pct_.*|trg_n|skor)$")


def sheet(wb, title, rows, widths=None, first=False):
    head = list(rows[0].keys())
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(head)
    for r in rows:
        ws.append([int(v) if NUM.match(h) and isinstance(v, str) and re.fullmatch(r"\d+", v.strip()) else v for h, v in r.items()])
    for c in ws[1]:
        c.font, c.fill, c.alignment = HF, HFILL, Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = NF
    for i, h in enumerate(head, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h) or min(60, max(10, len(h) + 4))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


with zipfile.ZipFile(f"{P}/Formulir_TRG_semua.zip", "w", zipfile.ZIP_DEFLATED) as z:
    for f in sorted(glob.glob("forms/trg/*.xlsx")):
        z.write(f, "Formulir_TRG/" + os.path.basename(f))
    z.write("forms/index_formulir.csv", "Formulir_TRG/_index_formulir.csv")
    z.write("forms/README.md", "Formulir_TRG/_BACA_DULU.md")

lap, idx = R("laporan_kelengkapan"), list(csv.DictReader(open("forms/index_formulir.csv", encoding="utf-8")))
wb = Workbook(); ws = wb.active; ws.title = "Ringkasan"
ws["A1"] = "Laporan kelengkapan data TRG Unhas"; ws["A1"].font = Font(name="Arial", size=14, bold=True)
ws["A2"] = 'Sumber: halaman publik trg.unhas.ac.id. Rumus menghitung dari lembar "Per TRG".'; ws["A2"].font = Font(name="Arial", size=9, italic=True)
n = len(lap) + 1
cnt = lambda k: f"=COUNTIF('Per TRG'!L2:L{n},\"*{k}*\")"
items = [("Jumlah TRG", f"=COUNTA('Per TRG'!C2:C{n})"), ("Total peneliti (baris, belum dedup)", f"=SUM('Per TRG'!E2:E{n})"),
         ("TRG tanpa masalah data", f"=COUNTBLANK('Per TRG'!L2:L{n})"), ("TRG keahlian <50%", cnt("keahlian_<50%")),
         ("TRG ID peneliti <30%", cnt("ID_peneliti_<30%")), ("TRG tanpa publikasi", cnt("tanpa_publikasi")),
         ("TRG tanpa proyek", cnt("tanpa_proyek")), ("TRG tim tidak terbaca", cnt("tim_tidak_terbaca")),
         ("TRG halaman tim 404", cnt("halaman_tim_404")), ("TRG belum punya halaman sendiri", cnt("tanpa_halaman_situs"))]
for i, (a, b) in enumerate(items, 4):
    ws[f"A{i}"], ws[f"B{i}"] = a, b
    ws[f"A{i}"].font = ws[f"B{i}"].font = NF
ws.column_dimensions["A"].width = 42; ws.column_dimensions["B"].width = 14
ws[f"A{len(items) + 6}"] = '"Tim tidak terbaca" berarti tata letak halaman TRG belum bisa dibaca otomatis, bukan berarti tim kosong.'
ws[f"A{len(items) + 6}"].font = Font(name="Arial", size=9, italic=True)
sheet(wb, "Per TRG", lap, {"nama": 52, "fakultas": 34, "masalah": 70, "slug": 18})
sheet(wb, "Prioritas formulir", idx, {"nama": 52, "fakultas": 34, "berkas": 42, "prioritas_isi": 90})
wb.save(f"{P}/Laporan_Kelengkapan_TRG.xlsx")

wb = Workbook(); first = True
for name, fn, w in [("Peneliti unik", "peneliti_unik", {"nama": 40, "trg_slugs": 40, "keahlian_bersih": 60, "nama_norm": 30}),
                    ("TRG", "trg", {"nama": 52, "deskripsi": 40, "fakultas": 34}), ("Tema Tingkat 1", "tag_tier1", {"definisi": 60, "nama_id": 44, "name_en": 40}),
                    ("Pemetaan TRG-Tema", "trg_tag_tier1", {"nama": 52, "fakultas": 34}), ("Usulan tag keahlian", "usulan_tag", {"varian": 70, "label_usulan": 36, "tag_key": 30}),
                    ("Kata kunci TRG", "trg_kata_kunci", {"nama": 52, "kata_kunci_dari_judul": 90}), ("Perlu dicek (dedup)", "peneliti_review", {"trg_slugs": 50})]:
    rows = R(fn)
    if fn == "peneliti_unik":
        rows = [{k: v for k, v in r.items() if k != "scholar_url"} for r in rows]
    sheet(wb, name, rows, w, first); first = False
lead = {}
for u in R("peneliti_unik"):
    for sl in [x for x in u["ketua_slugs"].split(";") if x]:
        lead.setdefault(sl, []).append(u)
dugaan = {(r["trg_slug"], r["nama"]) for r in csv.DictReader(open(f"{D}/ketua_diduga.csv", encoding="utf-8"))} if os.path.exists(f"{D}/ketua_diduga.csv") else set()
rows = []
for t in R("trg"):
    ls = lead.get(t["slug"], [])
    rows.append({"slug": t["slug"], "TRG": t["nama"], "fakultas": t["fakultas"],
                 "ketua": " | ".join(u["nama"] for u in ls) or "(belum terbaca dari situs)",
                 "dasar": ("diduga (periksa)" if any((t["slug"], u["nama"]) in dugaan or any((t["slug"], x) in dugaan for x in [u["nama"]]) for u in ls) else "berlabel Leader di situs") if ls else "",
                 "peneliti_id": " | ".join(u["peneliti_id"] for u in ls),
                 "juga_anggota_di": " | ".join(x for u in ls for x in u["trg_slugs"].split(";") if x and x != t["slug"]),
                 "sinta_id": " | ".join(u["sinta_id"] for u in ls if u["sinta_id"]), "scopus_id": " | ".join(u["scopus_id"] for u in ls if u["scopus_id"])})
sheet(wb, "Ketua TRG", rows, {"TRG": 52, "fakultas": 34, "ketua": 60, "juga_anggota_di": 40})
wb.save(f"{P}/Data_Peta_Riset.xlsx")
open(f"{P}/BACA_DULU.txt", "w", encoding="utf-8").write("""Paket Peta Riset Unhas (data publik trg.unhas.ac.id)

1. Formulir_TRG_semua.zip        Satu formulir xlsx per TRG, terisi dari data situs. Abu-abu = periksa, kuning = isi.
                                 _index_formulir.csv mengurutkan TRG yang paling perlu mengisi.
2. Laporan_Kelengkapan_TRG.xlsx  Ringkasan, masalah data per TRG, dan prioritas pengiriman formulir.
3. Data_Peta_Riset.xlsx          Peneliti unik, daftar TRG, tema Tingkat 1 (draf), pemetaan TRG-tema, usulan tag keahlian.

Tema Tingkat 1 dan pemetaannya adalah draf otomatis yang perlu ditinjau. Email peneliti tidak disertakan.
Dashboard: https://claude.ai/artifact/WFE6Tu1F7UNMKZFpwQQ7Nf
""")
print(sorted(os.listdir(P)))
